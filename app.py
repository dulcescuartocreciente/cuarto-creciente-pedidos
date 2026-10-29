import io
import hashlib
import hmac
import base64
import sqlite3
import secrets
import requests
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
import streamlit as st

BASE = Path(__file__).resolve().parent
# La base de datos permanente es Supabase; no se utiliza pedidos.db.
TZ = ZoneInfo('America/Argentina/Buenos_Aires')

st.set_page_config(page_title='Cuarto Creciente | Pedidos', page_icon='📦', layout='wide')
st.markdown('''<style>.block-container{padding-top:1.1rem;max-width:1250px} [data-testid="stMetric"]{background:#f2f6fa;padding:12px;border-radius:12px} </style>''', unsafe_allow_html=True)
def show_logo():
    for path in (BASE / 'logo.png', BASE / 'assets' / 'logo.png'):
        if path.is_file():
            st.image(str(path), width=480)
            return
    st.title('Cuarto Creciente')


def check_password(password, stored):
    try:
        method, rounds, salt64, hash64 = stored.split('$')
        if method != 'pbkdf2_sha256':
            return False
        derived = hashlib.pbkdf2_hmac('sha256', password.encode(), base64.b64decode(salt64), int(rounds))
        return hmac.compare_digest(derived, base64.b64decode(hash64))
    except (ValueError, TypeError):
        return False


def hash_password(password):
    salt = secrets.token_bytes(16)
    rounds = 600000
    derived = hashlib.pbkdf2_hmac('sha256', password.encode(), salt, rounds)
    return f"pbkdf2_sha256${rounds}${base64.b64encode(salt).decode()}${base64.b64encode(derived).decode()}"


def remote_config():
    try:
        url = str(st.secrets.get('SUPABASE_URL', '')).rstrip('/')
        key = str(st.secrets.get('SUPABASE_SERVICE_ROLE_KEY', ''))
    except (KeyError, FileNotFoundError):
        return None
    return (url, key) if url.startswith('https://') and key else None


def remote_users(method='GET', username=None, data=None):
    cfg = remote_config()
    if not cfg:
        raise RuntimeError('Falta configurar SUPABASE_URL y SUPABASE_SERVICE_ROLE_KEY en Streamlit Secrets.')
    url, key = cfg
    endpoint = f'{url}/rest/v1/app_users'
    params = {'select':'username,name,role,active,password_hash'} if method == 'GET' else {}
    if username is not None:
        params['username'] = f'eq.{username}'
    # Las claves nuevas sb_secret_ no son JWT: no enviarlas como Bearer.
    headers = {'apikey': key, 'Content-Type': 'application/json', 'Prefer': 'return=minimal'}
    if not key.startswith('sb_secret_'):
        headers['Authorization'] = f'Bearer {key}'
    try:
        response = requests.request(method, endpoint, params=params, json=data, headers=headers, timeout=12)
    except requests.RequestException as exc:
        raise RuntimeError('No se pudo conectar con Supabase. Revisá la conexión y la URL del proyecto.') from exc
    if not response.ok:
        # Mostrar solamente el código diagnóstico, nunca el cuerpo completo ni las claves.
        safe_code = 'sin código'
        try:
            payload = response.json()
            if isinstance(payload, dict):
                candidate = str(payload.get('code') or payload.get('error_code') or '')
                if candidate and len(candidate) <= 80 and all(ch.isalnum() or ch in '_-' for ch in candidate):
                    safe_code = candidate
        except ValueError:
            pass
        key_type = 'nueva sb_secret' if key.startswith('sb_secret_') else ('JWT heredada' if key.count('.') == 2 else 'formato desconocido')
        raise RuntimeError(f'Supabase respondió HTTP {response.status_code}. Código: {safe_code}. Tipo de clave: {key_type}. No compartas claves.')
    return response.json() if method == 'GET' else None


def login():
    if st.session_state.get('authenticated'):
        return True
    # Diseño compacto en dos columnas: imagen a la izquierda, formulario a la derecha.
    st.markdown("""<style>
      .block-container {padding-top: .55rem !important; padding-bottom: .5rem !important; max-width: 1100px !important;}
      [data-testid="stVerticalBlock"] {gap: .55rem;}
      [data-testid="stForm"] {padding: .75rem 1rem; border-radius: 14px;}
      [data-testid="stForm"] [data-testid="stVerticalBlock"] {gap: .45rem;}
      @media (max-width: 767px) {.block-container {padding-top: .35rem !important;}}
    </style>""", unsafe_allow_html=True)
    left, right = st.columns([1.12, 1], gap='large', vertical_alignment='center')
    with left:
        imagen_ingreso = BASE / 'imagen_ingreso.jpg'
        if imagen_ingreso.is_file():
            st.image(str(imagen_ingreso), use_container_width=True)
    with right:
        logo = BASE / 'logo.png'
        if not logo.is_file():
            logo = BASE / 'assets' / 'logo.png'
        if logo.is_file():
            st.image(str(logo), width=340)
        else:
            st.title('Cuarto Creciente')
        st.subheader('Acceso al sistema de pedidos')
        st.caption('Ingresá con tu usuario y contraseña')
        st.caption(datetime.now(TZ).strftime('Fecha: %d/%m/%Y · Hora: %H:%M (Argentina)'))
        try:
            users = st.secrets['users']
        except (KeyError, FileNotFoundError):
            st.error('Falta configurar los usuarios en los Secrets de Streamlit Cloud. Consultá el README.')
            return False
        with st.form('login_form'):
            username = st.text_input('Usuario').strip().lower()
            password = st.text_input('Contraseña', type='password')
            submitted = st.form_submit_button('Ingresar', use_container_width=True, type='primary')
        if submitted:
            user = None
            if username == 'admin':
                # La cuenta administradora local no depende de la API remota.
                user = users.get('admin')
            elif remote_config():
                try:
                    found = remote_users(username=username)
                    if found and found[0].get('active') and found[0].get('role') == 'vendedor':
                        user = found[0]
                except RuntimeError:
                    st.error('No se pudo consultar Supabase para ingresar como vendedor. El administrador puede ingresar normalmente.')
                    return False
            if user and check_password(password, user.get('password_hash', '')) and user.get('role') in ('admin', 'vendedor'):
                st.session_state.authenticated = True
                st.session_state.username = username
                st.session_state.role = user['role']
                st.session_state.display_name = user.get('name', username)
                st.rerun()
            else:
                st.error('Usuario o contraseña incorrectos.')
    return False


if not login():
    st.stop()

show_logo()
st.caption('Sistema de pedidos · Cuarto Creciente S.A.')
with st.sidebar:
    st.write(f"**{st.session_state.display_name}**")
    st.caption('Administrador' if st.session_state.role == 'admin' else 'Vendedor')
    st.link_button('📊 Seguimiento de pedidos (Drive)', 'https://docs.google.com/spreadsheets/d/1mRNTiXaUAF-lm5q0D1zr5zUPu7XrtN0FvH1u2lwVotw/edit?gid=0#gid=0', use_container_width=True)
    if st.button('Cerrar sesión'):
        st.session_state.clear()
        st.rerun()


# Supabase es la única fuente de datos de clientes, artículos y pedidos.
def api(method, path, params=None, data=None, prefer=None):
    cfg = remote_config()
    if not cfg:
        raise RuntimeError('Falta configurar Supabase en los Secrets de Streamlit.')
    url, key = cfg
    headers = {'apikey': key, 'Content-Type': 'application/json'}
    if not key.startswith('sb_secret_'):
        headers['Authorization'] = f'Bearer {key}'
    if prefer:
        headers['Prefer'] = prefer
    try:
        res = requests.request(method, f'{url}/rest/v1/{path}', headers=headers,
                               params=params, json=data, timeout=35)
    except requests.RequestException as exc:
        raise RuntimeError('No se pudo conectar con Supabase.') from exc
    if not res.ok:
        raise RuntimeError(f'Supabase: HTTP {res.status_code}. Revisá que el SQL se haya ejecutado y los permisos de la base.')
    return res.json() if res.content else None


def fetch_table(name):
    # Paginación para no perder filas si se supera el límite de la API.
    rows = []
    offset = 0
    while True:
        chunk = api('GET', name, params={'select': '*', 'limit': 1000, 'offset': offset}) or []
        rows.extend(chunk)
        if len(chunk) < 1000:
            break
        offset += 1000
    return rows


@st.cache_data(ttl=15, show_spinner=False)
def snapshot():
    return {name: fetch_table(name) for name in ('clientes', 'articulos', 'pedidos', 'renglones')}


def refresh():
    snapshot.clear()


def query(sql, params=()):
    # Consultas existentes de informes sobre una copia temporal en memoria.
    # NUNCA se guardan datos comerciales en el disco de Streamlit.
    data = snapshot()
    con = sqlite3.connect(':memory:')
    try:
        for name, rows in data.items():
            if name == 'clientes':
                schema = 'codigo TEXT, nombre TEXT, activo INTEGER, actualizado_en TEXT'
            elif name == 'articulos':
                schema = 'codigo TEXT, nombre TEXT, precio_1 REAL, precio_2 REAL, precio_3 REAL, activo INTEGER, actualizado_en TEXT'
            elif name == 'pedidos':
                schema = 'id INTEGER, fecha TEXT, cliente_codigo TEXT, cliente_nombre TEXT, lista TEXT, vendedor TEXT, pct_comun REAL, pct_x REAL, estado TEXT'
            else:
                schema = 'id INTEGER, pedido_id INTEGER, articulo_codigo TEXT, articulo_nombre TEXT, cantidad INTEGER, precio REAL'
            con.execute(f'CREATE TABLE {name} ({schema})')
            if rows:
                cols = [c[1] for c in con.execute(f'PRAGMA table_info({name})')]
                values = [tuple(row.get(c) for c in cols) for row in rows]
                con.executemany(f'INSERT INTO {name} ({",".join(cols)}) VALUES ({",".join("?" for _ in cols)})', values)
        return pd.read_sql_query(sql, con, params=params)
    finally:
        con.close()


def excel_bytes(sheets):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name[:31], index=False)
            ws = writer.sheets[name[:31]]
            ws.freeze_panes = 'A2'
            ws.auto_filter.ref = ws.dimensions
            from openpyxl.styles import Font, PatternFill
            for cell in ws[1]:
                cell.fill = PatternFill('solid', fgColor='214A70')
                cell.font = Font(color='FFFFFF', bold=True)
            for col in ws.columns:
                letter = col[0].column_letter
                max_len = max((len(str(c.value or '')) for c in list(col)[:300]), default=10)
                ws.column_dimensions[letter].width = min(max(14, max_len + 3), 44)
    return buf.getvalue()


def read_upload(file):
    if file.name.lower().endswith('.csv'):
        try:
            return pd.read_csv(file, dtype=str, sep=None, engine='python').fillna('')
        except UnicodeDecodeError:
            file.seek(0)
            return pd.read_csv(file, dtype=str, sep=None, engine='python', encoding='latin1').fillna('')
    return pd.read_excel(file, dtype=str).fillna('')


def import_catalog(kind, file):
    df = read_upload(file)
    df.columns = [str(x).strip().lower() for x in df.columns]
    required = ['codigo', 'nombre']
    if not set(required).issubset(df.columns):
        raise ValueError('Se necesitan columnas llamadas codigo y nombre.')
    if kind == 'clientes':
        # La lista se elige para cada pedido; no se guarda ni se importa en clientes.
        cols = ['codigo','nombre']
    else:
        for col in ['precio_1','precio_2','precio_3']:
            if col not in df: df[col] = '0'
            df[col] = pd.to_numeric(df[col].astype(str).str.replace(',','.', regex=False), errors='raise').fillna(0)
            if (df[col] < 0).any(): raise ValueError('Los precios no pueden ser negativos.')
        cols = ['codigo','nombre','precio_1','precio_2','precio_3']
    df = df[cols].copy()
    df['codigo'] = df['codigo'].astype(str).str.strip()
    df['nombre'] = df['nombre'].astype(str).str.strip()
    if (df['codigo']=='').any() or (df['nombre']=='').any() or df['codigo'].duplicated().any():
        raise ValueError('Hay códigos duplicados, códigos vacíos o nombres vacíos.')
    # Vista previa y actualización exclusiva de registros realmente cambiados.
    existentes = {r['codigo']: r for r in fetch_table(kind)}
    nuevos, modificados, sin_cambios = [], [], 0
    for row in df.to_dict('records'):
        row = {k: (float(v) if k.startswith('precio_') else str(v)) for k, v in row.items()}
        previo = existentes.get(row['codigo'])
        if previo is None:
            nuevos.append(row)
        elif any((round(float(previo[k]), 2) != round(float(row[k]), 2)) if k.startswith('precio_')
                 else str(previo[k]) != str(row[k]) for k in cols if k != 'codigo'):
            modificados.append(row)
        else:
            sin_cambios += 1
    return nuevos, modificados, sin_cambios


def apply_catalog(kind, nuevos, modificados, sin_cambios):
    # Upsert por código, únicamente para los nuevos o modificados.
    for group in (nuevos, modificados):
        for start in range(0, len(group), 200):
            api('POST', kind, data=group[start:start + 200],
                prefer='resolution=merge-duplicates,return=minimal')
    api('POST', 'importaciones', data={
        'usuario': st.session_state.username, 'catalogo': kind,
        'nuevos': len(nuevos), 'modificados': len(modificados),
        'sin_cambios': sin_cambios}, prefer='return=minimal')
    refresh()


if not remote_config():
    st.error('No está configurada la conexión a Supabase. No se pueden guardar pedidos.')
    st.stop()

now = datetime.now(TZ)
st.caption(f'Fecha y hora (Argentina): {now:%d/%m/%Y %H:%M:%S} · Se actualiza al interactuar o recargar la página')
def formulario_alta_cliente(origen):
    """Alta compartida para vendedores y administración, con validación de duplicados."""
    st.subheader('➕ Ingresar nuevo cliente')
    st.caption('El cliente se guarda en Supabase y queda disponible para todos los vendedores.')
    with st.form(f'alta_cliente_{origen}', clear_on_submit=True):
        nuevo_codigo = st.text_input('Código del cliente', key=f'codigo_cliente_{origen}',
                                     help='Usá el código del sistema comercial; se conservan los ceros iniciales.')
        nuevo_nombre = st.text_input('Nombre o razón social', key=f'nombre_cliente_{origen}')
        guardar_cliente = st.form_submit_button('Guardar nuevo cliente', type='primary')
    if guardar_cliente:
        codigo_limpio = nuevo_codigo.strip()
        nombre_limpio = nuevo_nombre.strip()
        if not codigo_limpio or not nombre_limpio:
            st.error('Completá el código y el nombre del cliente.')
        elif len(codigo_limpio) > 100 or len(nombre_limpio) > 250:
            st.error('El código o el nombre es demasiado largo.')
        else:
            try:
                existente = api('GET', 'clientes', params={'select': 'codigo,nombre', 'codigo': f'eq.{codigo_limpio}'}) or []
                if existente:
                    st.error(f'El código {codigo_limpio} ya pertenece a {existente[0]["nombre"]}. No se modificó.')
                else:
                    api('POST', 'clientes', data={'codigo': codigo_limpio, 'nombre': nombre_limpio}, prefer='return=minimal')
                    refresh()
                    st.session_state['cliente_recien_creado'] = codigo_limpio
                    st.toast(f'Cliente {nombre_limpio} guardado correctamente.', icon='✅')
                    st.rerun()
            except RuntimeError as exc:
                st.error(f'No se pudo guardar el cliente: {exc}')


def formulario_alta_articulo():
    """Carga manual de artículos, sin alterar la importación masiva por Excel."""
    st.subheader('➕ Agregar artículo manualmente')
    st.caption('Ingresá el código, la descripción y los precios por bulto de las tres listas.')
    with st.form('alta_articulo_manual', clear_on_submit=True):
        codigo = st.text_input('Código del artículo', help='Se guarda como texto, incluidos los ceros iniciales.')
        nombre = st.text_input('Nombre / descripción del artículo')
        c1, c2, c3 = st.columns(3)
        with c1:
            precio_1 = st.number_input('Precio lista 1 ($)', min_value=0.0, step=1.0, format='%.2f')
        with c2:
            precio_2 = st.number_input('Precio lista 2 ($)', min_value=0.0, step=1.0, format='%.2f')
        with c3:
            precio_3 = st.number_input('Precio lista 3 ($)', min_value=0.0, step=1.0, format='%.2f')
        guardar = st.form_submit_button('Guardar artículo', type='primary')
    if guardar:
        codigo, nombre = codigo.strip(), nombre.strip()
        if not codigo or not nombre:
            st.error('Completá el código y el nombre del artículo.')
        elif len(codigo) > 100 or len(nombre) > 250:
            st.error('El código o el nombre es demasiado largo.')
        else:
            try:
                existente = api('GET', 'articulos', params={'select': 'codigo,nombre', 'codigo': f'eq.{codigo}'}) or []
                if existente:
                    st.error(f'Ya existe el artículo {codigo}: {existente[0]["nombre"]}. No se modificó.')
                else:
                    api('POST', 'articulos', data={
                        'codigo': codigo, 'nombre': nombre,
                        'precio_1': float(precio_1), 'precio_2': float(precio_2),
                        'precio_3': float(precio_3)
                    }, prefer='return=minimal')
                    refresh()
                    st.session_state['articulo_creado_mensaje'] = f'Artículo {codigo} · {nombre} guardado correctamente.'
                    st.rerun()
            except RuntimeError as exc:
                st.error(f'No se pudo guardar el artículo: {exc}')


page = st.sidebar.radio('Menú', ['Nuevo pedido','Pedidos guardados','Exportar Excel','Administración','Usuarios'] if st.session_state.role == 'admin' else ['Nuevo pedido','Mis pedidos'])

if page == 'Nuevo pedido':
    st.header('Nuevo pedido')
    # Indicadores globales: un cliente se cuenta una sola vez aunque tenga varios pedidos.
    resumen_clientes = query('''SELECT
        (SELECT COUNT(*) FROM clientes) AS clientes_totales,
        (SELECT COUNT(DISTINCT cliente_codigo) FROM pedidos WHERE estado != 'Cancelado') AS clientes_con_pedidos''').iloc[0]
    total_clientes = int(resumen_clientes['clientes_totales'])
    clientes_con_pedidos = int(resumen_clientes['clientes_con_pedidos'])
    k1, k2, k3 = st.columns(3)
    k1.metric('Clientes totales', f'{total_clientes:,}'.replace(',', '.'))
    k2.metric('Clientes con pedidos', f'{clientes_con_pedidos:,}'.replace(',', '.'))
    k3.metric('Clientes sin pedidos', f'{max(0, total_clientes-clientes_con_pedidos):,}'.replace(',', '.'))
    st.caption('Clientes con pedidos: clientes distintos con al menos un pedido no cancelado, de todos los vendedores.')
    # Actividad comercial: cada cliente se cuenta una vez por vendedor.
    if st.session_state.role == 'admin':
        st.subheader('Clientes atendidos por vendedor')
        actividad = query("""SELECT vendedor AS Vendedor,
            COUNT(DISTINCT cliente_codigo) AS 'Clientes con pedidos',
            COUNT(*) AS 'Pedidos realizados'
            FROM pedidos WHERE estado != 'Cancelado'
            GROUP BY vendedor ORDER BY 2 DESC, 3 DESC, vendedor""")
        if actividad.empty:
            st.info('Todavía no hay pedidos registrados por vendedores.')
        else:
            st.dataframe(actividad, use_container_width=True, hide_index=True)
            st.caption('Un cliente atendido por dos vendedores se cuenta una vez para cada uno; en el total general se cuenta una sola vez.')
            st.download_button('⬇️ Descargar actividad por vendedor',
                data=excel_bytes({'Actividad vendedores': actividad}),
                file_name=f'actividad_vendedores_{datetime.now(TZ):%Y%m%d}.xlsx',
                mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    else:
        propios = query("""SELECT COUNT(DISTINCT cliente_codigo) AS clientes,
            COUNT(*) AS pedidos FROM pedidos
            WHERE estado != 'Cancelado' AND vendedor = ?""", (st.session_state.username,)).iloc[0]
        a1, a2 = st.columns(2)
        a1.metric('Mis clientes atendidos', int(propios['clientes']))
        a2.metric('Mis pedidos realizados', int(propios['pedidos']))
    with st.expander('➕ ¿El cliente no existe? Dar de alta un cliente nuevo'):
        formulario_alta_cliente('vendedor' if st.session_state.role == 'vendedor' else 'nuevo_pedido')
    clientes = query('SELECT * FROM clientes ORDER BY nombre')
    articulos = query('SELECT * FROM articulos ORDER BY nombre')
    if clientes.empty or articulos.empty:
        st.warning('Primero importá los clientes y artículos desde Administración. No se cargaron datos ficticios.')
    else:
        lookup_c = {r.codigo: r for r in clientes.itertuples(index=False)}
        lookup_a = {r.codigo: r for r in articulos.itertuples(index=False)}
        c1,c2 = st.columns([2,1])
        with c1:
            codigo_preseleccionado = st.session_state.pop('cliente_recien_creado', None)
            opciones_clientes = clientes['codigo'].tolist()
            indice_cliente = opciones_clientes.index(codigo_preseleccionado) if codigo_preseleccionado in opciones_clientes else 0
            cliente = st.selectbox('Cliente: escribí código o nombre', opciones_clientes, index=indice_cliente, format_func=lambda c: f'{c} · {lookup_c[c].nombre}')
        with c2:
            vendedor = st.session_state.username
            st.text_input('Vendedor', value=st.session_state.display_name, disabled=True)
        c = lookup_c[cliente]
        st.info(f'Cliente: {c.nombre} · Código: {cliente}')
        lista = st.selectbox(
            'Lista de precios para este pedido',
            ['Lista 1', 'Lista 2', 'Lista 3'],
            help='El vendedor selecciona la lista según el volumen negociado. Se guarda con el pedido.'
        )
        st.caption('Lista 1: general · Lista 2: más de 300 bultos · Lista 3: pallet completo / línea de producción. Verificá las condiciones comerciales antes de guardar.')
        st.subheader('Facturación')
        if 'pct_comun' not in st.session_state: st.session_state.pct_comun = 100.0
        if 'pct_x' not in st.session_state: st.session_state.pct_x = 0.0
        def update_comun():
            v = st.session_state.pct_comun
            if 0 <= v <= 100: st.session_state.pct_x = round(100-v,2)
        def update_x():
            v = st.session_state.pct_x
            if 0 <= v <= 100: st.session_state.pct_comun = round(100-v,2)
        a,b = st.columns(2)
        a.number_input('Factura común (%)', min_value=0.0, max_value=100.0, step=5.0, key='pct_comun', on_change=update_comun)
        b.number_input('Factura X (%)', min_value=0.0, max_value=100.0, step=5.0, key='pct_x', on_change=update_x)
        st.caption('Ambos porcentajes se guardan como información administrativa. No se generan comprobantes fiscales.')
        st.subheader('Artículos')
        st.caption('Escribí los bultos al lado de cada artículo. Dejá 0 para no incluirlo en el pedido.')
        if 'carrito' not in st.session_state:
            st.session_state.carrito = {}
        # Reiniciar cantidades al comienzo de la ejecución siguiente, antes de crear widgets.
        if st.session_state.pop('reiniciar_cantidades_pedido', False):
            for code in lookup_a:
                st.session_state[f'cantidad_art_{code}'] = 0
        # Una casilla por código; el valor se conserva al buscar y al cambiar la lista.
        for code in lookup_a:
            widget_key = f'cantidad_art_{code}'
            if widget_key not in st.session_state:
                st.session_state[widget_key] = int(st.session_state.carrito.get(code, 0))
        buscar_articulo = st.text_input('🔎 Buscar por código o nombre', placeholder='Ej.: JBP1400 o mermelada', key='buscar_lista_articulos')
        filtro = buscar_articulo.strip().casefold()
        visibles = [r for r in articulos.itertuples(index=False) if filtro in str(r.codigo).casefold() or filtro in str(r.nombre).casefold()]
        st.caption(f'Mostrando {len(visibles)} de {len(articulos)} artículos')
        with st.container(height=440, border=True):
            for art in visibles:
                col_art, col_precio, col_cant = st.columns([5, 2, 2], vertical_alignment='center')
                with col_art:
                    st.markdown(f'**{art.codigo}** · {art.nombre}')
                with col_precio:
                    precio_art = float(getattr(art, 'precio_' + lista[-1]))
                    st.caption(f'$ {precio_art:,.2f} / bulto')
                with col_cant:
                    st.number_input(f'Bultos · {art.codigo}', min_value=0, max_value=100000, step=1, key=f'cantidad_art_{art.codigo}', label_visibility='collapsed')
        rows = []
        for code, item in lookup_a.items():
            qty = int(st.session_state.get(f'cantidad_art_{code}', 0))
            if qty > 0:
                price = float(getattr(item, 'precio_' + lista[-1]))
                rows.append({'Código artículo':code,'Artículo':item.nombre,'Bultos':qty,'Precio/bulto':price,'Importe':qty*price})
        total_bultos = sum(r['Bultos'] for r in rows)
        total_importe = sum(r['Importe'] for r in rows)
        pct_x = float(st.session_state.pct_x)
        pct_comun = float(st.session_state.pct_comun)
        st.subheader('Resumen del pedido')
        f_x, f_comun, f_total = st.columns(3)
        f_x.metric('Factura X', f'$ {total_importe * pct_x / 100:,.2f}')
        f_comun.metric('Factura común', f'$ {total_importe * pct_comun / 100:,.2f}')
        f_total.metric('TOTAL', f'$ {total_importe:,.2f}')
        st.metric('Total de bultos', total_bultos)
        if rows:
            with st.expander(f'Ver detalle de los {len(rows)} artículos seleccionados'):
                st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)
        else:
            st.info('Todavía no ingresaste cantidades. Escribí los bultos junto a los artículos.')
        if st.button('💾 Guardar pedido', type='primary', disabled=not rows):
            if not vendedor.strip(): st.error('Indicá el vendedor.')
            elif round(st.session_state.pct_comun+st.session_state.pct_x,2)!=100:
                st.error('Los porcentajes deben sumar 100%.')
            else:
                pid = api('POST', 'rpc/crear_pedido', data={
                    'p_cliente_codigo': cliente, 'p_cliente_nombre': c.nombre,
                    'p_lista': lista, 'p_vendedor': vendedor.strip(),
                    'p_pct_comun': float(st.session_state.pct_comun),
                    'p_pct_x': float(st.session_state.pct_x),
                    'p_renglones': [{'articulo_codigo': r['Código artículo'],
                                    'articulo_nombre': r['Artículo'],
                                    'cantidad': r['Bultos'], 'precio': r['Precio/bulto']}
                                   for r in rows]})
                refresh()
                st.session_state.carrito={}
                st.session_state.reiniciar_cantidades_pedido = True
                st.success(f'Pedido #{pid} guardado correctamente. Para iniciar otro pedido, actualizá la pantalla.')
                st.caption('No se descuenta stock: la asignación se realizará en Distribución Inteligente.')

elif page in ('Pedidos guardados', 'Mis pedidos'):
    st.header('Pedidos guardados')
    pedidos=query('SELECT p.id,p.fecha,p.cliente_codigo,p.cliente_nombre,p.lista,p.vendedor,p.pct_comun,p.pct_x,p.estado,COALESCE(SUM(r.cantidad),0) AS bultos FROM pedidos p LEFT JOIN renglones r ON r.pedido_id=p.id WHERE (? = \'admin\' OR p.vendedor = ?) GROUP BY p.id ORDER BY p.id DESC', (st.session_state.role, st.session_state.username))
    if st.session_state.role == 'admin' and not pedidos.empty:
        st.subheader('Gestión masiva de estados')
        st.caption('Seleccioná pedidos con las casillas y aplicá un estado a todos los seleccionados. Los cambios se guardan en Supabase.')
        opciones = ['Todos', 'Pendiente', 'Revisado', 'Exportado', 'Cancelado']
        filtro_estado = st.selectbox('Filtrar por estado', opciones, key='filtro_estado_pedidos')
        vista = pedidos if filtro_estado == 'Todos' else pedidos[pedidos['estado'] == filtro_estado]
        if not vista.empty:
            ids_visibles = [int(v) for v in vista['id'].tolist()]
            if st.button('Seleccionar todos los pedidos visibles', key='seleccionar_pedidos_visibles'):
                for numero in ids_visibles:
                    st.session_state[f'marcar_pedido_{numero}'] = True
                st.rerun()
            seleccion = vista[['id', 'fecha', 'cliente_codigo', 'cliente_nombre', 'vendedor', 'bultos', 'estado']].copy()
            seleccion.insert(0, 'Seleccionar', [bool(st.session_state.get(f'marcar_pedido_{int(numero)}', False)) for numero in seleccion['id']])
            # El editor usa una clave propia por filtro; se guardan las selecciones entre filtros.
            editado = st.data_editor(seleccion, hide_index=True, use_container_width=True,
                disabled=[col for col in seleccion.columns if col != 'Seleccionar'],
                column_config={'Seleccionar': st.column_config.CheckboxColumn('✓', help='Marcar pedido para cambiar estado')},
                key=f'tabla_estados_{filtro_estado}')
            for numero, marcado in zip(editado['id'], editado['Seleccionar']):
                st.session_state[f'marcar_pedido_{int(numero)}'] = bool(marcado)
            elegidos = [int(numero) for numero in pedidos['id'] if st.session_state.get(f'marcar_pedido_{int(numero)}', False)]
            st.write(f'**{len(elegidos)} pedidos seleccionados**')
            with st.form('cambio_masivo_estados'):
                nuevo_estado = st.selectbox('Editar estado', ['Pendiente', 'Revisado', 'Exportado', 'Cancelado'])
                aplicar = st.form_submit_button('Aplicar estado a los seleccionados', type='primary', disabled=not elegidos)
            if aplicar and elegidos:
                try:
                    # Se aplica únicamente a los IDs seleccionados, no a los demás pedidos.
                    for inicio in range(0, len(elegidos), 100):
                        lote = elegidos[inicio:inicio + 100]
                        api('PATCH', 'pedidos', params={'id': f'in.({",".join(map(str, lote))})'}, data={'estado': nuevo_estado})
                    refresh()
                    for numero in elegidos:
                        st.session_state.pop(f'marcar_pedido_{numero}', None)
                    st.success(f'{len(elegidos)} pedidos actualizados a {nuevo_estado}.')
                    st.rerun()
                except RuntimeError as exc:
                    st.error(f'No se pudo completar el cambio de estados: {exc}')
        else:
            st.info('No hay pedidos con ese estado.')
        st.divider()
        st.subheader('Detalle de pedidos')
    else:
        st.dataframe(pedidos, use_container_width=True, hide_index=True)
    if not pedidos.empty:
        pid=st.selectbox('Ver o editar pedido', pedidos['id'].tolist())
        detalle=query('SELECT articulo_codigo,articulo_nombre,cantidad,precio FROM renglones WHERE pedido_id=?',(int(pid),))
        st.dataframe(detalle, hide_index=True, use_container_width=True)
        if st.session_state.role == 'admin':
          with st.expander('Editar cantidades / estado'):
            cambios={}
            for r in detalle.itertuples(index=False):
                cambios[r.articulo_codigo]=st.number_input(f'{r.articulo_codigo} · {r.articulo_nombre}',min_value=0,value=int(r.cantidad),step=1,key=f'edit_{pid}_{r.articulo_codigo}')
            estado_actual=pedidos.loc[pedidos['id']==pid,'estado'].iloc[0]
            estados=['Pendiente','Revisado','Exportado','Cancelado']
            estado=st.selectbox('Estado',estados,index=estados.index(estado_actual) if estado_actual in estados else 0)
            if st.button('Guardar cambios'):
                for code, qty in cambios.items():
                    filters = {'pedido_id': f'eq.{int(pid)}', 'articulo_codigo': f'eq.{code}'}
                    if qty == 0:
                        api('DELETE', 'renglones', params=filters)
                    else:
                        api('PATCH', 'renglones', params=filters, data={'cantidad': int(qty)})
                api('PATCH', 'pedidos', params={'id': f'eq.{int(pid)}'}, data={'estado': estado})
                refresh()
                st.success('Pedido actualizado.')
                st.rerun()

elif page == 'Exportar Excel':
    st.header('Exportación para Distribución Inteligente')
    st.caption('Se exporta la demanda total, sin separar los artículos por tipo de facturación.')
    df=query('''SELECT p.id AS pedido,p.fecha,p.cliente_codigo,p.cliente_nombre,p.lista,p.vendedor,p.pct_comun,p.pct_x,p.estado,r.articulo_codigo,r.articulo_nombre,r.cantidad AS bultos,r.precio,(r.cantidad*r.precio) AS importe FROM pedidos p JOIN renglones r ON p.id=r.pedido_id ORDER BY p.id,r.articulo_codigo''')
    if df.empty: st.info('No hay pedidos para exportar todavía.')
    else:
        opciones_estado = sorted(df.estado.dropna().unique().tolist())
        estados = st.multiselect('Estados incluidos', opciones_estado, default=[e for e in ['Pendiente', 'Revisado'] if e in opciones_estado])
        dff=df[df.estado.isin(estados)].copy()
        if dff.empty: st.warning('No hay renglones con los estados seleccionados.')
        else:
            matriz=dff.pivot_table(index=['cliente_codigo','cliente_nombre'],columns='articulo_codigo',values='bultos',aggfunc='sum',fill_value=0).reset_index()
            matriz.columns.name=None
            totales=dff.groupby(['articulo_codigo','articulo_nombre'],as_index=False)['bultos'].sum()
            resumen=dff.groupby(['cliente_codigo','cliente_nombre'],as_index=False).agg(bultos=('bultos','sum'),importe=('importe','sum'))
            st.dataframe(dff,use_container_width=True,hide_index=True)
            data=excel_bytes({'Detalle pedidos':dff,'Matriz distribucion':matriz,'Totales articulos':totales,'Resumen clientes':resumen})
            st.download_button('⬇️ Descargar Excel de pedidos',data=data,file_name=f'pedidos_cuarto_creciente_{datetime.now(TZ):%Y%m%d_%H%M}.xlsx',mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',type='primary')
            st.divider()
            st.subheader('📦 Excel para Distribuidor')
            st.caption('Archivo independiente: CODIGO, DESCRIPCION, CANTIDAD y CLIENTE. Suma bultos del mismo artículo por cliente.')
            expedicion = dff[['cliente_codigo', 'cliente_nombre', 'articulo_codigo', 'articulo_nombre', 'bultos']].copy()
            expedicion['cliente_codigo'] = expedicion['cliente_codigo'].astype(str)
            expedicion['articulo_codigo'] = expedicion['articulo_codigo'].astype(str)
            expedicion['bultos'] = pd.to_numeric(expedicion['bultos'], errors='raise').astype(int)
            # Agrupar por códigos, para evitar mezclar clientes con el mismo nombre.
            expedicion = (expedicion.groupby(['cliente_codigo', 'articulo_codigo'], as_index=False, sort=False)
                          .agg({'cliente_nombre': 'first', 'articulo_nombre': 'first', 'bultos': 'sum'}))
            expedicion = expedicion.rename(columns={
                'articulo_codigo': 'CODIGO', 'articulo_nombre': 'DESCRIPCION',
                'bultos': 'CANTIDAD', 'cliente_nombre': 'CLIENTE'})
            expedicion = expedicion[['CODIGO', 'DESCRIPCION', 'CANTIDAD', 'CLIENTE']]
            expedicion = expedicion.sort_values(['CLIENTE', 'CODIGO'], kind='stable').reset_index(drop=True)
            st.dataframe(expedicion, use_container_width=True, hide_index=True)
            st.download_button('⬇️ Descargar Excel para Distribuidor',
                data=excel_bytes({'Distribuidor': expedicion}),
                file_name=f'distribuidor_cuarto_creciente_{datetime.now(TZ):%Y%m%d_%H%M}.xlsx',
                mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                type='primary', key='descargar_expedicion')
            st.caption('La descarga completa se mantiene separada. Los pedidos cancelados solo se incluyen si los seleccionás.')

elif page == 'Usuarios':
    st.header('Administración de usuarios')
    st.caption('Solo administradores pueden crear, desactivar y cambiar contraseñas de vendedores.')
    if not remote_config():
        st.warning('Para activar esta sección primero hay que conectar una base compartida. Seguí las instrucciones de INSTALACION_USUARIOS.md. El administrador actual sigue entrando por Secrets.')
    else:
        try:
            registered = remote_users()
            st.dataframe(pd.DataFrame([{'Usuario':u['username'],'Nombre':u['name'],'Rol':u['role'],'Activo':u['active']} for u in registered]), hide_index=True, use_container_width=True)
            with st.form('create_user', clear_on_submit=True):
                st.subheader('Crear vendedor')
                new_name = st.text_input('Nombre y apellido')
                new_username = st.text_input('Usuario (sin espacios)')
                new_password = st.text_input('Contraseña inicial', type='password')
                create = st.form_submit_button('Crear vendedor', type='primary')
            if create:
                un = new_username.strip().lower()
                if not un or not un.replace('_','').isalnum() or not new_name.strip():
                    st.error('Completá nombre y usuario. Usá letras, números o guion bajo.')
                elif len(new_password) < 12:
                    st.error('La contraseña debe tener al menos 12 caracteres.')
                elif any(u['username'] == un for u in registered) or un == 'admin':
                    st.error('Ese usuario ya existe.')
                else:
                    remote_users('POST', data={'username':un,'name':new_name.strip(),'role':'vendedor','active':True,'password_hash':hash_password(new_password)})
                    st.success('Vendedor creado. Ya puede ingresar desde la pantalla principal.')
                    st.rerun()
            st.divider()
            sellers = [u for u in registered if u['role'] == 'vendedor']
            if sellers:
                chosen = st.selectbox('Administrar vendedor', [u['username'] for u in sellers])
                current = next(u for u in sellers if u['username'] == chosen)
                if st.button('Desactivar usuario' if current['active'] else 'Reactivar usuario'):
                    remote_users('PATCH', chosen, {'active':not current['active']})
                    st.rerun()
                with st.form('reset_password'):
                    new_pass = st.text_input('Nueva contraseña para este vendedor',type='password')
                    change = st.form_submit_button('Cambiar contraseña')
                if change:
                    if len(new_pass)<12:
                        st.error('La contraseña debe tener al menos 12 caracteres.')
                    else:
                        remote_users('PATCH', chosen, {'password_hash':hash_password(new_pass)})
                        st.success('Contraseña actualizada.')
        except RuntimeError as exc:
            st.error(str(exc))

else:
    st.header('Administración de catálogos')
    st.info('Clientes, artículos y pedidos se guardan en Supabase. Importar Excel actualiza solo los registros con diferencias.')
    formulario_alta_cliente('administracion')
    st.divider()
    if st.session_state.get('articulo_creado_mensaje'):
        st.success(st.session_state.pop('articulo_creado_mensaje'))
    formulario_alta_articulo()
    st.divider()
    st.subheader('Importación por Excel (opción existente)')
    for kind, title, cols in [('clientes','Clientes',['codigo','nombre']),('articulos','Artículos',['codigo','nombre','precio_1','precio_2','precio_3'])]:
        st.subheader(title)
        current=query(f'SELECT codigo,nombre FROM clientes ORDER BY codigo' if kind == 'clientes' else 'SELECT * FROM articulos ORDER BY codigo')
        st.dataframe(current,use_container_width=True,hide_index=True)
        example=pd.DataFrame(columns=cols)
        st.download_button(f'Descargar plantilla de {kind}',excel_bytes({title:example}),file_name=f'plantilla_{kind}.xlsx',key=f'template_{kind}')
        file=st.file_uploader(f'Importar {title.lower()} (Excel o CSV)',type=['xlsx','csv'],key=f'upload_{kind}')
        if file:
            try:
                nuevos, modificados, sin_cambios = import_catalog(kind, file)
                st.write(f'**Vista previa:** {len(nuevos)} nuevos · {len(modificados)} modificados · {sin_cambios} sin cambios.')
                if modificados:
                    st.dataframe(pd.DataFrame(modificados).head(30), hide_index=True, use_container_width=True)
                if st.button(f'Confirmar actualización de {title.lower()}', key=f'btn_{kind}',
                             disabled=not (nuevos or modificados)):
                    apply_catalog(kind, nuevos, modificados, sin_cambios)
                    st.success('Importación aplicada en Supabase.')
                    st.rerun()
            except (ValueError, RuntimeError) as e:
                st.error(f'No se pudo importar: {e}')
    st.divider()
    st.subheader('Copia de seguridad')
    st.download_button('⬇️ Exportar clientes y artículos',
        data=excel_bytes({'Clientes': query('SELECT codigo,nombre FROM clientes ORDER BY codigo'),
                          'Artículos': query('SELECT codigo,nombre,precio_1,precio_2,precio_3 FROM articulos ORDER BY codigo')}),
        file_name='catalogos_cuarto_creciente.xlsx',
        mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    st.caption('Los códigos deben conservarse como texto para no perder ceros iniciales. Importar actualiza registros existentes por código sin borrar los demás.')
